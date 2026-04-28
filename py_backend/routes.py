from __future__ import annotations

import hashlib
import io
import math
import re
from datetime import datetime, timedelta
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import JSONResponse, Response, StreamingResponse
from openpyxl import Workbook, load_workbook

from py_backend.normalization import normalize_article
from py_backend.context import reload_smart_cache
from py_backend.storage import DatabaseStorage, InsufficientBoxStockError, InsufficientStockError, InvalidRequestError
from py_backend.types import SaleStatus

MAX_IMPORT_FILE_BYTES = 10 * 1024 * 1024
DEFAULT_MEDIA_CHUNK_SIZE = 1024 * 1024

# Example: "bytes=0-99" / "bytes=100-" / "bytes=-500"
_RANGE_RE = re.compile(r"^bytes=(\d*)-(\d*)$")
BUSINESS_TZ = ZoneInfo("Europe/Moscow")


def _parse_range_header(range_header: str, size_bytes: int) -> tuple[int, int] | None:
    if size_bytes <= 0:
        return None
    header = (range_header or "").strip()
    if not header:
        return None
    m = _RANGE_RE.match(header)
    if not m:
        return None

    start_s, end_s = m.group(1), m.group(2)
    if start_s == "" and end_s == "":
        return None

    if start_s == "":
        # Suffix range: last N bytes
        suffix = int(end_s)
        if suffix <= 0:
            return None
        start = max(0, size_bytes - suffix)
        end = size_bytes - 1
        return (start, end)

    start = int(start_s)
    if start < 0:
        return None
    if end_s == "":
        end = size_bytes - 1
    else:
        end = int(end_s)
        if end < start:
            return None
        end = min(end, size_bytes - 1)
    if start >= size_bytes:
        return None
    return (start, end)


def parse_csv_text(text: str) -> list[list[str]]:
    rows: list[list[str]] = []
    row: list[str] = []
    field = ""
    in_quotes = False

    def push_field() -> None:
        nonlocal field
        row.append(field)
        field = ""

    def push_row() -> None:
        nonlocal row
        if len(row) == 1 and row[0] == "" and len(rows) > 0:
            row = []
            return
        rows.append(row)
        row = []

    i = 0
    while i < len(text):
        ch = text[i]

        if in_quotes:
            if ch == '"':
                nxt = text[i + 1] if i + 1 < len(text) else None
                if nxt == '"':
                    field += '"'
                    i += 1
                else:
                    in_quotes = False
            else:
                field += ch
            i += 1
            continue

        if ch == '"':
            in_quotes = True
            i += 1
            continue

        if ch == ",":
            push_field()
            i += 1
            continue

        if ch == "\n":
            push_field()
            push_row()
            i += 1
            continue

        if ch == "\r":
            nxt = text[i + 1] if i + 1 < len(text) else None
            if nxt == "\n":
                i += 1
            else:
                push_field()
                push_row()
            i += 1
            continue

        field += ch
        i += 1

    push_field()
    if len(row) > 0:
        push_row()

    while len(rows) > 0 and all(cell == "" for cell in rows[-1]):
        rows.pop()
    return rows


def pick_row_value(obj: dict[str, Any], keys: list[str]) -> Any:
    for key in keys:
        if key in obj and obj[key] is not None:
            return obj[key]
    return None


def to_string_or_empty(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        return str(int(value)) if value.is_integer() else str(value)
    return ""


def to_number_or_zero(value: Any) -> float:
    if isinstance(value, bool):
        return 0
    if isinstance(value, (int, float)):
        n = float(value)
        return n if math.isfinite(n) else 0
    raw = to_string_or_empty(value)
    try:
        n = float(raw)
        return n if math.isfinite(n) else 0
    except Exception:
        return 0


def to_int_or_none_strict(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return None
        return int(value)
    raw = to_string_or_empty(value).strip()
    if not raw:
        return None
    try:
        n = float(raw)
    except Exception:
        return None
    if not math.isfinite(n) or not n.is_integer():
        return None
    return int(n)


def to_opt_string(value: Any) -> str | None:
    text = to_string_or_empty(value).strip()
    return text or None


async def parse_json_body(request: Request) -> dict[str, Any]:
    try:
        body = await request.json()
    except Exception:
        raise InvalidRequestError("Некорректный JSON") from None
    if not isinstance(body, dict):
        raise InvalidRequestError("JSON body должен быть объектом")
    return body


def _today_bounds_moscow_naive() -> tuple[datetime, datetime]:
    now = datetime.now(BUSINESS_TZ)
    start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    end = start + timedelta(days=1)
    return (start.replace(tzinfo=None), end.replace(tzinfo=None))


def parse_bulk_import_rows_from_objects(raw_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for idx, obj in enumerate(raw_rows):
        smart = to_string_or_empty(pick_row_value(obj, ["smart", "SMART"])).strip()
        qty_delta = to_number_or_zero(pick_row_value(obj, ["qty_delta", "qtyDelta", "qty"]))
        reason = to_string_or_empty(pick_row_value(obj, ["reason", "type"])).strip()

        row: dict[str, Any] = {"smart": smart, "qtyDelta": qty_delta, "reason": reason}

        note = to_opt_string(pick_row_value(obj, ["note", "comment", "примечание"]))
        if note:
            row["note"] = note

        purchase_price = to_opt_string(pick_row_value(obj, ["purchase_price", "purchasePrice"]))
        if purchase_price:
            row["purchasePrice"] = purchase_price

        sale_price = to_opt_string(pick_row_value(obj, ["sale_price", "salePrice"]))
        if sale_price:
            row["salePrice"] = sale_price

        delivery_price = to_opt_string(pick_row_value(obj, ["delivery_price", "deliveryPrice"]))
        if delivery_price:
            row["deliveryPrice"] = delivery_price

        box_number = to_opt_string(pick_row_value(obj, ["box_number", "boxNumber"]))
        if box_number:
            row["boxNumber"] = box_number

        track_number = to_opt_string(pick_row_value(obj, ["track_number", "trackNumber"]))
        if track_number:
            row["trackNumber"] = track_number

        shipping_method_id_raw = pick_row_value(obj, ["shipping_method_id", "shippingMethodId"])
        shipping_method_id_num = to_int_or_none_strict(shipping_method_id_raw)
        if shipping_method_id_num is not None and shipping_method_id_num > 0:
            row["shippingMethodId"] = shipping_method_id_num

        row["__row"] = obj.get("__row", idx + 2)
        result.append(row)

    return result


def _storage_from_request(request: Request) -> DatabaseStorage:
    return request.app.state.ctx.storage


def register_routes(app: FastAPI) -> None:
    @app.get("/api/articles/search")
    async def search_articles(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            query = request.query_params.get("query")
            if not query:
                return JSONResponse(status_code=400, content={"error": "Query parameter is required"})

            normalized = normalize_article(query)
            if not normalized or len(normalized) < 2:
                return JSONResponse(status_code=400, content={"error": "Query parameter is too short"})

            limit_raw = request.query_params.get("limit")
            limit = 50
            if limit_raw:
                try:
                    limit = max(1, min(50, int(limit_raw)))
                except (ValueError, TypeError):
                    pass

            matches = storage.searchSmart(normalized, limit=limit)
            smart_codes = [m["smart"] for m in matches]
            stock_map = await storage.getTotalStockBySmartBatch(smart_codes)
            payload = [{**m, "currentStock": stock_map.get(m["smart"], 0)} for m in matches]
            return JSONResponse(content=payload)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to search"})

    @app.get("/api/smart/{code}")
    async def get_smart(request: Request, code: str) -> Response:
        storage = _storage_from_request(request)
        smart = storage.getSmartByCode(code)
        if not smart:
            return JSONResponse(status_code=404, content={"error": "SMART code not found"})
        if hasattr(smart, "model_dump"):
            return JSONResponse(content=smart.model_dump())
        return JSONResponse(content=smart)

    @app.get("/api/customers")
    async def get_customers(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            query = request.query_params.get("query")
            include_archived = request.query_params.get("includeArchived") in {"1", "true"}
            customers = await storage.getCustomers(query, include_archived)
            return JSONResponse(content=customers)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get customers"})

    @app.get("/api/customers/{customer_id}")
    async def get_customer_by_id(request: Request, customer_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            cid = int(customer_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            customer = await storage.getCustomerById(cid)
            if not customer:
                return JSONResponse(status_code=404, content={"error": "Customer not found"})

            orders = await storage.getOrders({"customerId": cid, "includeArchivedCustomers": True})
            total_amount = 0.0
            total_returns = 0
            for order in orders:
                order_id = order.get("id")
                if not isinstance(order_id, int):
                    continue
                details = await storage.getOrderById(order_id)
                if not details:
                    continue
                items_total = sum(
                    max(0, int(item.get("qty", 0)) - int(item.get("returnedQty", 0))) * float(item.get("salePrice") or 0)
                    for item in details.get("items", [])
                )
                buyer_delivery = sum(
                    float(shipment.get("deliveryPrice") or 0)
                    for shipment in details.get("shipments", [])
                    if shipment.get("deliveryPayer") == "buyer"
                )
                total_amount += items_total + buyer_delivery
                total_returns += len(details.get("returns", []))
            stats = {
                "ordersCount": len(orders),
                "totalAmount": round(total_amount, 2),
                "returnsCount": total_returns,
            }
            return JSONResponse(content={"customer": customer, "orders": orders, "stats": stats})
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get customer"})

    @app.post("/api/customers")
    async def create_customer(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            body = await parse_json_body(request)
            customer = await storage.createCustomer(body)
            return JSONResponse(status_code=201, content=customer)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to create customer"})

    @app.patch("/api/customers/{customer_id}")
    async def update_customer(request: Request, customer_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            cid = int(customer_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            body = await parse_json_body(request)
            customer = await storage.updateCustomer(cid, body)
            return JSONResponse(content=customer)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            if str(err) == "Customer not found":
                return JSONResponse(status_code=404, content={"error": str(err)})
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to update customer"})

    @app.post("/api/orders")
    async def create_order(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            body = await parse_json_body(request)
            order = await storage.createOrder(body)
            return JSONResponse(status_code=201, content=order)
        except InsufficientBoxStockError as err:
            return JSONResponse(
                status_code=409,
                content={
                    "error": str(err),
                    "details": {
                        "smart": err.smart,
                        "boxName": err.boxName,
                        "available": err.available,
                        "requested": err.requested,
                    },
                },
            )
        except InsufficientStockError as err:
            return JSONResponse(
                status_code=409,
                content={
                    "error": str(err),
                    "details": {"smart": err.smart, "currentStock": err.currentStock, "requestedQty": err.requestedQty},
                },
            )
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to create order"})

    @app.get("/api/orders")
    async def get_orders(request: Request) -> Response:
        storage = _storage_from_request(request)
        customer_id_raw = request.query_params.get("customerId")
        customer_id: int | None = None
        if customer_id_raw and customer_id_raw.strip():
            try:
                customer_id = int(customer_id_raw)
            except Exception:
                return JSONResponse(status_code=400, content={"error": "Invalid customerId"})

        try:
            include_archived = request.query_params.get("includeArchived") in {"1", "true"}
            orders = await storage.getOrders(
                {
                    "customerId": customer_id,
                    "includeArchivedCustomers": include_archived,
                }
            )
            return JSONResponse(content=orders)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get orders"})

    @app.get("/api/orders/{order_id}")
    async def get_order_by_id(request: Request, order_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            oid = int(order_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            order = await storage.getOrderById(oid)
            if not order:
                return JSONResponse(status_code=404, content={"error": "Order not found"})
            return JSONResponse(content=order)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get order details"})

    @app.post("/api/orders/{order_id}/shipments")
    async def create_additional_shipment(request: Request, order_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            oid = int(order_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            body = await parse_json_body(request)
            shipment = await storage.createAdditionalShipment(oid, body)
            return JSONResponse(status_code=201, content=shipment)
        except InvalidRequestError as err:
            message = str(err)
            if message == "Заказ не найден":
                return JSONResponse(status_code=404, content={"error": message})
            return JSONResponse(status_code=400, content={"error": message})
        except Exception as err:
            return JSONResponse(
                status_code=500,
                content={"error": str(err) if str(err) else "Failed to create shipment"},
            )

    @app.patch("/api/shipments/{shipment_id}/status")
    async def update_shipment_status(request: Request, shipment_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            sid = int(shipment_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            body = await parse_json_body(request)
            status = body.get("status")
            if status not in ("pending", "shipped", "delivered"):
                return JSONResponse(status_code=400, content={"error": "Input should be 'pending', 'shipped' or 'delivered'"})
            shipment = await storage.updateShipmentStatus(sid, status)
            return JSONResponse(content=shipment)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            if str(err) == "Shipment not found":
                return JSONResponse(status_code=404, content={"error": str(err)})
            return JSONResponse(
                status_code=500,
                content={"error": str(err) if str(err) else "Failed to update shipment status"},
            )

    @app.post("/api/orders/{order_id}/returns")
    async def create_order_return(request: Request, order_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            oid = int(order_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            body = await parse_json_body(request)
            created = await storage.createOrderReturn(oid, body)
            return JSONResponse(status_code=201, content=created)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to create order return"})

    @app.post("/api/movements")
    async def create_movement(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            body = await parse_json_body(request)
            if body.get("reason") == "return":
                return JSONResponse(
                    status_code=400,
                    content={"error": "Возврат создается только через страницу проданных товаров"},
                )
            movement = await storage.createMovement(body)
            return JSONResponse(status_code=201, content=movement)
        except InsufficientBoxStockError as err:
            return JSONResponse(
                status_code=409,
                content={
                    "error": str(err),
                    "details": {
                        "smart": err.smart,
                        "boxName": err.boxName,
                        "available": err.available,
                        "requested": err.requested,
                    },
                },
            )
        except InsufficientStockError as err:
            return JSONResponse(
                status_code=409,
                content={
                    "error": str(err),
                    "details": {"smart": err.smart, "currentStock": err.currentStock, "requestedQty": err.requestedQty},
                },
            )
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to create movement"})

    @app.post("/api/movements/batch")
    async def create_movements_batch(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            body = await parse_json_body(request)
            items = body.get("items")
            if not isinstance(items, list) or len(items) == 0:
                return JSONResponse(status_code=400, content={"error": "items должен быть непустым массивом"})
            for item in items:
                if item.get("reason") == "return":
                    return JSONResponse(
                        status_code=400,
                        content={"error": "Возврат создается только через страницу проданных товаров"},
                    )
            results = await storage.createMovementsBatch(items)
            return JSONResponse(status_code=201, content=results)
        except InsufficientBoxStockError as err:
            return JSONResponse(
                status_code=409,
                content={
                    "error": str(err),
                    "details": {
                        "smart": err.smart,
                        "boxName": err.boxName,
                        "available": err.available,
                        "requested": err.requested,
                    },
                },
            )
        except InsufficientStockError as err:
            return JSONResponse(
                status_code=409,
                content={
                    "error": str(err),
                    "details": {"smart": err.smart, "currentStock": err.currentStock, "requestedQty": err.requestedQty},
                },
            )
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to create movements batch"})

    @app.get("/api/movements")
    async def get_movements(request: Request) -> Response:
        storage = _storage_from_request(request)
        box_number = request.query_params.get("boxNumber")
        try:
            movements = await storage.getMovements({"boxNumber": box_number} if box_number else None)
            return JSONResponse(content=movements)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get movements"})

    @app.get("/api/movements/{movement_id}/items")
    async def get_movement_items(request: Request, movement_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            mid = int(movement_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            items = await storage.getMovementItems(mid)
            return JSONResponse(content=items)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get movement items"})

    @app.get("/api/items")
    async def get_items(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            data = await storage.getItems(
                {
                    "smart": request.query_params.get("smart"),
                    "boxNumber": request.query_params.get("boxNumber"),
                    "state": request.query_params.get("state"),
                    "q": request.query_params.get("q"),
                    "limit": request.query_params.get("limit"),
                    "offset": request.query_params.get("offset"),
                }
            )
            return JSONResponse(content=data)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to get items"})

    @app.get("/api/items/{item_id}")
    async def get_item_by_id(request: Request, item_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            iid = int(item_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})
        try:
            item = await storage.getItemById(iid)
            if item is None:
                return JSONResponse(status_code=404, content={"error": "Item not found"})
            return JSONResponse(content=item)
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to get item"})

    @app.patch("/api/items/{item_id}")
    async def update_item(request: Request, item_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            iid = int(item_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            body = await parse_json_body(request)
            updates: dict[str, Any] = {}
            if "note" in body:
                updates["note"] = body.get("note")
            updated = await storage.updateItem(iid, updates)
            return JSONResponse(content=updated)
        except InvalidRequestError as err:
            msg = str(err)
            if msg == "Item not found":
                return JSONResponse(status_code=404, content={"error": msg})
            return JSONResponse(status_code=400, content={"error": msg})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to update item"})

    @app.post("/api/items/{item_id}/media")
    async def upload_item_media(
        request: Request,
        item_id: str,
        file: UploadFile | None = File(default=None),
        kind: str | None = Form(default=None),
    ) -> Response:
        storage = _storage_from_request(request)
        try:
            iid = int(item_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})
        if file is None:
            return JSONResponse(status_code=400, content={"error": "No file uploaded"})

        mime = (file.content_type or "").strip() or "application/octet-stream"
        inferred_kind = "video" if mime.lower().startswith("video/") else "photo"
        actual_kind = (kind or inferred_kind).strip().lower()
        if actual_kind not in ("photo", "video"):
            return JSONResponse(status_code=400, content={"error": "Invalid kind. Must be 'photo' or 'video'."})

        filename = (file.filename or "").strip() or None

        conn = await storage.inventoryPool.acquire()
        try:
            await conn.execute("BEGIN")
            try:
                exists = await conn.fetchrow("SELECT id FROM inventory.items WHERE id = $1", iid)
                if exists is None:
                    await conn.execute("ROLLBACK")
                    return JSONResponse(status_code=404, content={"error": "Item not found"})

                media_row = await conn.fetchrow(
                    """
                    INSERT INTO inventory.item_media (
                      item_id, kind, filename, mime, size_bytes, sha256, chunk_size, created_at
                    )
                    VALUES ($1, $2, $3, $4, 0, NULL, $5, NOW())
                    RETURNING id, chunk_size
                    """,
                    iid,
                    actual_kind,
                    filename,
                    mime,
                    DEFAULT_MEDIA_CHUNK_SIZE,
                )
                if media_row is None:
                    raise Exception("Failed to create media record")
                media_id = int(media_row.get("id"))  # type: ignore[call-arg]
                chunk_size = int(media_row.get("chunk_size") or DEFAULT_MEDIA_CHUNK_SIZE)  # type: ignore[call-arg]

                hasher = hashlib.sha256()
                total = 0
                idx = 0
                while True:
                    chunk = await file.read(chunk_size)
                    if not chunk:
                        break
                    total += len(chunk)
                    hasher.update(chunk)
                    await conn.execute(
                        """
                        INSERT INTO inventory.item_media_chunks (media_id, idx, data, created_at)
                        VALUES ($1, $2, $3, NOW())
                        """,
                        media_id,
                        idx,
                        chunk,
                    )
                    idx += 1

                sha256 = hasher.hexdigest() if total > 0 else None
                await conn.execute(
                    "UPDATE inventory.item_media SET size_bytes = $1, sha256 = $2 WHERE id = $3",
                    total,
                    sha256,
                    media_id,
                )

                await conn.execute("COMMIT")
                return JSONResponse(
                    status_code=201,
                    content={
                        "id": media_id,
                        "itemId": iid,
                        "kind": actual_kind,
                        "filename": filename,
                        "mime": mime,
                        "sizeBytes": total,
                        "sha256": sha256,
                        "chunkSize": chunk_size,
                    },
                )
            except Exception:
                await conn.execute("ROLLBACK")
                raise
        finally:
            await storage.inventoryPool.release(conn)

    @app.get("/api/item-media/{media_id}")
    async def get_item_media(request: Request, media_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            mid = int(media_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        conn = await storage.inventoryPool.acquire()
        try:
            meta = await conn.fetchrow(
                """
                SELECT id, filename, mime, size_bytes, chunk_size
                FROM inventory.item_media
                WHERE id = $1
                  AND deleted_at IS NULL
                """,
                mid,
            )
            if meta is None:
                await storage.inventoryPool.release(conn)
                return JSONResponse(status_code=404, content={"error": "Media not found"})

            filename = meta.get("filename")  # type: ignore[call-arg]
            mime = str(meta.get("mime") or "application/octet-stream")  # type: ignore[call-arg]
            size_bytes = int(meta.get("size_bytes") or 0)  # type: ignore[call-arg]
            chunk_size = int(meta.get("chunk_size") or DEFAULT_MEDIA_CHUNK_SIZE)  # type: ignore[call-arg]

            range_header = request.headers.get("range") or request.headers.get("Range") or ""
            parsed = _parse_range_header(range_header, size_bytes)
            is_partial = parsed is not None
            if parsed is None:
                start = 0
                end = max(0, size_bytes - 1)
            else:
                start, end = parsed

            if size_bytes == 0:
                await storage.inventoryPool.release(conn)
                return Response(status_code=200, content=b"", media_type=mime)

            start_idx = start // chunk_size
            end_idx = end // chunk_size
            start_off = start - (start_idx * chunk_size)
            end_off = end - (end_idx * chunk_size)
            content_length = end - start + 1

            async def gen():
                try:
                    for idx in range(start_idx, end_idx + 1):
                        row = await conn.fetchrow(
                            "SELECT data FROM inventory.item_media_chunks WHERE media_id = $1 AND idx = $2",
                            mid,
                            idx,
                        )
                        if row is None:
                            break
                        data = row.get("data")  # type: ignore[call-arg]
                        if not isinstance(data, (bytes, bytearray)):
                            continue
                        chunk_data = bytes(data)
                        if idx == start_idx:
                            chunk_data = chunk_data[start_off:]
                        if idx == end_idx:
                            chunk_data = chunk_data[: end_off + 1]
                        if chunk_data:
                            yield chunk_data
                finally:
                    await storage.inventoryPool.release(conn)

            headers: dict[str, str] = {
                "Accept-Ranges": "bytes",
                "Content-Length": str(content_length),
            }
            if is_partial:
                headers["Content-Range"] = f"bytes {start}-{end}/{size_bytes}"
            if filename:
                headers["Content-Disposition"] = f'inline; filename="{filename}"'

            return StreamingResponse(
                gen(),
                status_code=206 if is_partial else 200,
                media_type=mime,
                headers=headers,
            )
        except Exception:
            await storage.inventoryPool.release(conn)
            raise

    @app.delete("/api/item-media/{media_id}")
    async def delete_item_media(request: Request, media_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            mid = int(media_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            row = await storage.inventoryPool.fetchrow(
                "UPDATE inventory.item_media SET deleted_at = NOW() WHERE id = $1 AND deleted_at IS NULL RETURNING id",
                mid,
            )
            if row is None:
                return JSONResponse(status_code=404, content={"error": "Media not found"})
            return Response(status_code=204)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to delete media"})

    @app.get("/api/stock")
    async def get_stock(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            stock = await storage.getStockLevels()
            return JSONResponse(content=stock)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get stock levels"})

    @app.get("/api/stock/{smart}/purchases")
    async def get_stock_purchases(request: Request, smart: str) -> Response:
        storage = _storage_from_request(request)
        try:
            purchases = await storage.getPurchasesBySmart(smart)
            return JSONResponse(content=purchases)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get purchases"})

    @app.get("/api/stock/{smart}/sales")
    async def get_stock_sales(request: Request, smart: str) -> Response:
        storage = _storage_from_request(request)
        try:
            data = await storage.getSalesAnalyticsBySmart(smart)
            return JSONResponse(content=data)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get sales analytics"})

    @app.get("/api/stock/{smart}")
    async def get_stock_by_smart(request: Request, smart: str) -> Response:
        storage = _storage_from_request(request)
        try:
            info = await storage.getStockBySmart(smart)
            if not info.get("existed"):
                return JSONResponse(status_code=404, content={"error": "SMART code not found in inventory history"})
            # The frontend expects `existed` in the payload (shared/schema.ts).
            # Keeping it also makes the API response self-descriptive.
            return JSONResponse(content=info)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get stock details"})

    @app.patch("/api/movements/{movement_id}")
    async def update_movement(request: Request, movement_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            mid = int(movement_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            body = await parse_json_body(request)
            updates: dict[str, Any] = {}
            if "boxNumber" in body or "qtyDelta" in body:
                return JSONResponse(
                    status_code=400,
                    content={
                        "error": "Редактирование количества/коробки через PATCH запрещено в item-системе. "
                        "Используйте перемещение/корректировку."
                    },
                )
            if "purchasePrice" in body:
                updates["purchasePrice"] = body.get("purchasePrice")
            if "note" in body:
                updates["note"] = body.get("note")

            movement = await storage.updateMovement(mid, updates)
            return JSONResponse(content=movement)
        except InsufficientStockError as err:
            return JSONResponse(
                status_code=409,
                content={
                    "error": str(err),
                    "details": {"smart": err.smart, "currentStock": err.currentStock, "requestedQty": err.requestedQty},
                },
            )
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            if str(err) == "Movement not found":
                return JSONResponse(status_code=404, content={"error": str(err)})
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to update movement"})

    # Boxes (storage locations)
    @app.get("/api/boxes")
    async def get_boxes(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            active_only = request.query_params.get("activeOnly") in {"1", "true"}
            data = await storage.getBoxes(activeOnly=active_only)
            return JSONResponse(content=data)
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to get boxes"})

    @app.get("/api/unboxed")
    async def get_unboxed(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            items = await storage.getUnboxedItems()
            return JSONResponse(content=items)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get unboxed items"})

    @app.post("/api/boxes")
    async def create_box(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            body = await parse_json_body(request)
            created = await storage.createBox(body)
            return JSONResponse(status_code=201, content=created)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to create box"})

    @app.post("/api/boxes/transfer")
    async def transfer_between_boxes(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            body = await parse_json_body(request)
            created = await storage.transferBetweenBoxes(body)
            return JSONResponse(status_code=201, content=created)
        except InsufficientBoxStockError as err:
            return JSONResponse(
                status_code=409,
                content={
                    "error": str(err),
                    "details": {
                        "smart": err.smart,
                        "boxName": err.boxName,
                        "available": err.available,
                        "requested": err.requested,
                    },
                },
            )
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to transfer"})

    @app.get("/api/boxes/{box_name}")
    async def get_box_details(request: Request, box_name: str) -> Response:
        storage = _storage_from_request(request)
        try:
            limit_raw = request.query_params.get("historyLimit")
            limit = 50
            if limit_raw and limit_raw.strip():
                try:
                    limit = int(limit_raw)
                except Exception:
                    limit = 50
            details = await storage.getBoxDetails(box_name, historyLimit=limit)
            return JSONResponse(content=details)
        except InvalidRequestError as err:
            return JSONResponse(status_code=404, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to get box details"})

    @app.patch("/api/boxes/{box_name}")
    async def update_box(request: Request, box_name: str) -> Response:
        storage = _storage_from_request(request)
        try:
            body = await parse_json_body(request)
            updated = await storage.updateBox(box_name, body)
            return JSONResponse(content=updated)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to update box"})

    # Boxes containing a specific SMART (used for sales/writeoffs)
    @app.get("/api/stock/{smart}/boxes")
    async def get_smart_boxes(request: Request, smart: str) -> Response:
        storage = _storage_from_request(request)
        try:
            data = await storage.getBoxesForSmart(smart)
            return JSONResponse(content=data)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to get smart boxes"})

    @app.get("/api/reasons")
    async def get_reasons(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            return JSONResponse(content=storage.getReasons())
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get reasons"})

    @app.get("/api/shipping-methods")
    async def get_shipping_methods(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            methods = await storage.getShippingMethods()
            return JSONResponse(content=methods)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get shipping methods"})

    @app.post("/api/shipping-methods")
    async def create_shipping_method(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            body = await parse_json_body(request)
            name = body.get("name")
            if not name or not isinstance(name, str):
                return JSONResponse(status_code=400, content={"error": "Name is required"})
            method = await storage.createShippingMethod({"name": name, "isPickup": bool(body.get("isPickup"))})
            return JSONResponse(status_code=201, content=method)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            if getattr(err, "sqlstate", None) == "23505" or getattr(err, "code", None) == "23505":
                return JSONResponse(status_code=400, content={"error": "Такой способ доставки уже существует"})
            return JSONResponse(status_code=500, content={"error": "Failed to create shipping method"})

    @app.delete("/api/shipping-methods/{method_id}")
    async def delete_shipping_method(request: Request, method_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            mid = int(method_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})
        try:
            await storage.deleteShippingMethod(mid)
            return Response(status_code=204)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to delete shipping method"})

    @app.patch("/api/movements/{movement_id}/status")
    async def update_movement_status(request: Request, movement_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            mid = int(movement_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            body = await parse_json_body(request)
            status = body.get("status")
            if status not in ("awaiting_shipment", "shipped"):
                return JSONResponse(status_code=400, content={"error": "Invalid status"})
            parsed: SaleStatus = status
            movement = await storage.updateMovementSaleStatus(mid, parsed)
            return JSONResponse(content=movement)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            if str(err) == "Movement not found":
                return JSONResponse(status_code=404, content={"error": str(err)})
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to update movement status"})

    @app.patch("/api/movements/{movement_id}/ship")
    async def mark_movement_shipped(request: Request, movement_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            mid = int(movement_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})
        try:
            movement = await storage.updateMovementSaleStatus(mid, "shipped")
            return JSONResponse(content=movement)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            if str(err) == "Movement not found":
                return JSONResponse(status_code=404, content={"error": str(err)})
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to mark as shipped"})

    @app.get("/api/sold-out")
    async def get_sold_out(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            items = await storage.getSoldOutItems()
            return JSONResponse(content=items)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get sold out items"})

    @app.get("/api/top-parts")
    async def get_top_parts(request: Request) -> Response:
        storage = _storage_from_request(request)
        mode = request.query_params.get("mode")
        if mode not in ("profit", "sales", "combined"):
            return JSONResponse(status_code=400, content={"error": "Invalid mode. Must be 'profit', 'sales', or 'combined'"})
        try:
            items = await storage.getTopParts(mode)
            return JSONResponse(content=items)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get top parts"})

    @app.post("/api/movements/{movement_id}/return")
    async def return_movement(request: Request, movement_id: str) -> Response:
        storage = _storage_from_request(request)
        try:
            mid = int(movement_id)
        except Exception:
            return JSONResponse(status_code=400, content={"error": "Invalid ID"})

        try:
            body = await parse_json_body(request)
            box_number = body.get("boxNumber")
            if not box_number:
                return JSONResponse(status_code=400, content={"error": "Номер коробки обязателен"})
            note_raw = body.get("note")
            note = note_raw.strip() if isinstance(note_raw, str) and note_raw.strip() else None
            return_movement = await storage.returnLegacySaleMovement(mid, str(box_number), note)
            return JSONResponse(status_code=201, content=return_movement)
        except InvalidRequestError as err:
            msg = str(err)
            if msg in {"Продажа не найдена"}:
                return JSONResponse(status_code=404, content={"error": msg})
            if msg in {"Товар уже возвращен на склад"}:
                return JSONResponse(status_code=409, content={"error": msg})
            return JSONResponse(status_code=400, content={"error": msg})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to return to inventory"})

    @app.post("/api/bulk-import")
    async def bulk_import(request: Request, file: UploadFile | None = File(default=None)) -> Response:
        storage = _storage_from_request(request)
        try:
            if file is None:
                return JSONResponse(status_code=400, content={"error": "No file uploaded"})

            content = await file.read()
            if len(content) > MAX_IMPORT_FILE_BYTES:
                return JSONResponse(status_code=413, content={"error": f"File too large. Max {MAX_IMPORT_FILE_BYTES} bytes."})

            filename = (file.filename or "").lower()
            content_type = (file.content_type or "").lower()
            is_excel = "sheet" in content_type or filename.endswith(".xlsx") or filename.endswith(".xls")
            is_csv = "csv" in content_type or filename.endswith(".csv")

            rows: list[dict[str, Any]] = []

            if is_excel:
                if filename.endswith(".xls"):
                    try:
                        import xlrd  # type: ignore
                    except Exception:
                        return JSONResponse(status_code=400, content={"error": "Unsupported file type"})

                    book = xlrd.open_workbook(file_contents=content)
                    if book.nsheets == 0:
                        return JSONResponse(status_code=400, content={"error": "Empty Excel file (no sheets)"})
                    sheet = book.sheet_by_index(0)
                    if sheet.nrows == 0:
                        return JSONResponse(status_code=400, content={"error": "Empty Excel sheet"})
                    headers = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
                    raw_rows: list[dict[str, Any]] = []
                    for r in range(1, sheet.nrows):
                        values = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
                        if all(not str(v).strip() for v in values):
                            continue
                        obj: dict[str, Any] = {"__row": r + 1}
                        for c, h in enumerate(headers):
                            obj[h] = values[c] if c < len(values) else ""
                        raw_rows.append(obj)
                    rows = parse_bulk_import_rows_from_objects(raw_rows)
                else:
                    workbook = load_workbook(filename=io.BytesIO(content), data_only=True)
                    if not workbook.sheetnames:
                        return JSONResponse(status_code=400, content={"error": "Empty Excel file (no sheets)"})
                    worksheet = workbook[workbook.sheetnames[0]]
                    sheet_rows = list(worksheet.iter_rows(values_only=True))
                    if len(sheet_rows) == 0:
                        return JSONResponse(status_code=400, content={"error": "Empty Excel sheet"})

                    headers = [str(v).strip() if v is not None else "" for v in sheet_rows[0]]
                    raw_rows = []
                    for idx, values in enumerate(sheet_rows[1:], start=2):
                        vals = list(values)
                        if all(v is None or str(v).strip() == "" for v in vals):
                            continue
                        obj = {"__row": idx}
                        for col, header in enumerate(headers):
                            obj[header] = vals[col] if col < len(vals) and vals[col] is not None else ""
                        raw_rows.append(obj)
                    if len(raw_rows) == 0:
                        return JSONResponse(status_code=400, content={"error": "Empty Excel sheet"})
                    rows = parse_bulk_import_rows_from_objects(raw_rows)
            elif is_csv:
                csv_text = content.decode("utf-8")
                table = parse_csv_text(csv_text)
                if len(table) == 0:
                    return JSONResponse(status_code=400, content={"error": "Empty CSV file"})

                headers = [h.strip() for h in table[0]]
                raw_objects: list[dict[str, Any]] = []
                for r in range(1, len(table)):
                    values = table[r]
                    if all(not str(v).strip() for v in values):
                        continue
                    obj: dict[str, Any] = {"__row": r + 1}
                    for c, header in enumerate(headers):
                        obj[header] = values[c] if c < len(values) else ""
                    raw_objects.append(obj)
                rows = parse_bulk_import_rows_from_objects(raw_objects)
            else:
                return JSONResponse(status_code=400, content={"error": "Unsupported file type"})

            normalized_rows = []
            for row in rows:
                normalized_rows.append(
                    {
                        **row,
                        "smart": row.get("smart").strip() if isinstance(row.get("smart"), str) else row.get("smart"),
                        "reason": row.get("reason").strip() if isinstance(row.get("reason"), str) else row.get("reason"),
                    }
                )

            result = await storage.processBulkImport(normalized_rows)
            return JSONResponse(content=result)
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to process bulk import"})

    @app.get("/api/import-template")
    async def import_template(request: Request) -> Response:
        storage = _storage_from_request(request)
        shipping_method_id = 1
        try:
            methods = await storage.getShippingMethods()
            if methods and methods[0].get("id"):
                shipping_method_id = int(methods[0]["id"])
        except Exception:
            shipping_method_id = 1

        headers = [
            "smart",
            "qty_delta",
            "reason",
            "purchase_price",
            "sale_price",
            "delivery_price",
            "shipping_method_id",
            "box_number",
            "track_number",
            "note",
        ]

        # Keep examples importable with the current validation rules (box_number is required for all new operations).
        template_data = [
            {
                "smart": "smart_17713",
                "qty_delta": 10,
                "reason": "purchase",
                "purchase_price": 100.0,
                "box_number": "K-123",
                "note": "Example purchase",
            },
            {
                "smart": "smart_17713",
                "qty_delta": -1,
                "reason": "sale",
                "sale_price": 250.0,
                "delivery_price": 0.0,
                "shipping_method_id": shipping_method_id,
                "box_number": "K-123",
                "note": "Example sale (set shipping_method_id to an existing method)",
            },
            {
                "smart": "smart_17713",
                "qty_delta": -1,
                "reason": "writeoff",
                "box_number": "K-123",
                "note": "Example writeoff",
            },
            {
                "smart": "smart_17713",
                "qty_delta": 2,
                "reason": "adjust",
                "purchase_price": 120.0,
                "box_number": "K-123",
                "note": "Пересчет склада, нашли лишние",
            },
        ]

        wb = Workbook()
        ws = wb.active
        ws.title = "Import Template"
        ws.append(headers)
        for row in template_data:
            ws.append([row.get(h) for h in headers])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return Response(
            content=output.read(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="inventory-import-template.xlsx"'},
        )

    # ── SMART Catalog CRUD ─────────────────────────────────────────

    @app.get("/api/smart-catalog")
    async def list_smart_catalog(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            search = request.query_params.get("search")
            limit = min(500, max(1, int(request.query_params.get("limit", "200"))))
            offset = max(0, int(request.query_params.get("offset", "0")))
            result = await storage.getSmartCatalog(search=search, limit=limit, offset=offset)
            return JSONResponse(content=result)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to list smart catalog"})

    @app.post("/api/smart-catalog")
    async def create_smart_entry(request: Request) -> Response:
        storage = _storage_from_request(request)
        try:
            body = await parse_json_body(request)
            entry = await storage.createSmartEntry(body)
            ctx = request.app.state.ctx
            await reload_smart_cache(ctx)
            return JSONResponse(status_code=201, content=entry)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to create smart entry"})

    @app.put("/api/smart-catalog/{code:path}")
    async def update_smart_entry(request: Request) -> Response:
        storage = _storage_from_request(request)
        code = request.path_params["code"]
        try:
            body = await parse_json_body(request)
            entry = await storage.updateSmartEntry(code, body)
            ctx = request.app.state.ctx
            await reload_smart_cache(ctx)
            return JSONResponse(content=entry)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to update smart entry"})

    @app.delete("/api/smart-catalog/{code:path}")
    async def delete_smart_entry(request: Request) -> Response:
        storage = _storage_from_request(request)
        code = request.path_params["code"]
        try:
            result = await storage.deleteSmartEntry(code)
            ctx = request.app.state.ctx
            await reload_smart_cache(ctx)
            return JSONResponse(content=result)
        except InvalidRequestError as err:
            return JSONResponse(status_code=400, content={"error": str(err)})
        except Exception as err:
            return JSONResponse(status_code=500, content={"error": str(err) if str(err) else "Failed to delete smart entry"})

    @app.get("/api/dashboard/stats")
    async def dashboard_stats(request: Request) -> Response:
        ctx = request.app.state.ctx
        try:
            start_moscow, end_moscow = _today_bounds_moscow_naive()
            row = await ctx.pools.inventory_pool.fetchrow(
                """
                WITH
                  in_stock AS (
                    SELECT COUNT(*)::text as in_stock
                    FROM inventory.stock
                  ),
                  total_parts AS (
                    SELECT COALESCE(SUM(total_qty::bigint), 0)::text as total_parts
                    FROM inventory.stock
                  ),
                  movements_today AS (
                    SELECT
                      COUNT(*)::text as movements_today,
                      COALESCE(SUM(CASE WHEN reason = 'sale' THEN ABS(qty_delta) ELSE 0 END), 0)::text as sales_today
                    FROM inventory.movements
                    WHERE created_at >= $1
                      AND created_at < $2
                  )
                SELECT
                  (SELECT in_stock FROM in_stock) as in_stock,
                  (SELECT total_parts FROM total_parts) as total_parts,
                  (SELECT movements_today FROM movements_today) as movements_today,
                  (SELECT sales_today FROM movements_today) as sales_today
                """
                ,
                start_moscow,
                end_moscow,
            )
            return JSONResponse(
                content={
                    "inStock": int(row["in_stock"] if row else 0),
                    "totalParts": int(row["total_parts"] if row else 0),
                    "movementsToday": int(row["movements_today"] if row else 0),
                    "salesToday": int(row["sales_today"] if row else 0),
                }
            )
        except Exception:
            return JSONResponse(status_code=500, content={"error": "Failed to get dashboard stats"})

    @app.get("/api/health/db")
    async def db_health(request: Request) -> Response:
        ctx = request.app.state.ctx
        try:
            await ctx.pools.inventory_pool.fetchrow("SELECT 1")
            return JSONResponse(content={"connected": True})
        except Exception:
            return JSONResponse(status_code=503, content={"connected": False})
